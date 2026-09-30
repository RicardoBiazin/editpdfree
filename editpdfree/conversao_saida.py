"""Converter PDF para PowerPoint, Excel e Markdown, e extrair as imagens.

Cada formato tem o seu limite honesto:

* **PowerPoint** -- um slide por pagina: a pagina inteira como imagem de fundo
  (fica identica ao PDF) e, por cima, caixas de texto TRANSPARENTES com o texto
  de cada bloco, na mesma posicao. Da' para editar ou copiar o texto; o
  desenho de fundo nao se edita.
* **Excel** -- so' o que o detector de tabelas do PyMuPDF (`find_tables`)
  reconhece como tabela, uma aba por tabela. Numeros no formato brasileiro
  ("1.234,56") viram numero de verdade na planilha.
* **Markdown** -- titulos deduzidos pelo tamanho da fonte em relacao ao corpo
  do texto, negrito/italico, listas e tabelas.
"""

from __future__ import annotations

import collections
import io
import os
import pathlib
import re

import pymupdf

from .documento import Documento

PT_EMU = 12700          # EMU por ponto (unidade do Office)


# -- PowerPoint ---------------------------------------------------------------
def para_powerpoint(d: Documento, destino: str | os.PathLike, *,
                    dpi: int = 150, progresso=None) -> pathlib.Path:
    from pptx import Presentation
    from pptx.util import Emu, Pt

    apresentacao = Presentation()
    primeira = d.doc[0].rect
    apresentacao.slide_width = Emu(int(primeira.width * PT_EMU))
    apresentacao.slide_height = Emu(int(primeira.height * PT_EMU))
    em_branco = apresentacao.slide_layouts[6]
    for i in range(d.paginas):
        if progresso is not None and progresso(i, d.paginas) is False:
            break
        p = d.doc[i]
        slide = apresentacao.slides.add_slide(em_branco)
        # Pagina de tamanho diferente da primeira: escala para o slide.
        esc_x = primeira.width / p.rect.width
        esc_y = primeira.height / p.rect.height
        png = p.get_pixmap(dpi=dpi, alpha=False, annots=True).tobytes("png")
        slide.shapes.add_picture(io.BytesIO(png), 0, 0,
                                 width=apresentacao.slide_width,
                                 height=apresentacao.slide_height)
        # Texto por cima, no espaco que o leitor ve (pagina girada inclusive).
        matriz = p.rotation_matrix
        for bloco in p.get_text("dict")["blocks"]:
            if bloco["type"] != 0:
                continue
            texto = "\n".join("".join(s["text"] for s in l["spans"])
                              for l in bloco["lines"]).strip()
            if not texto:
                continue
            r = pymupdf.Rect(bloco["bbox"]) * matriz
            r.normalize()
            caixa = slide.shapes.add_textbox(
                Emu(int(r.x0 * esc_x * PT_EMU)), Emu(int(r.y0 * esc_y * PT_EMU)),
                Emu(int(max(r.width, 10) * esc_x * PT_EMU)),
                Emu(int(max(r.height, 10) * esc_y * PT_EMU)))
            quadro = caixa.text_frame
            quadro.word_wrap = True
            for margem in ("margin_left", "margin_right", "margin_top",
                           "margin_bottom"):
                setattr(quadro, margem, 0)
            tamanho = max(s["size"] for l in bloco["lines"] for s in l["spans"])
            quadro.text = texto
            for paragrafo in quadro.paragraphs:
                for run in paragrafo.runs:
                    run.font.size = Pt(max(1, tamanho * esc_y))
                    # Transparente: a imagem de fundo ja' mostra o texto; a
                    # caixa existe para selecionar, copiar e editar.
                    run.font.fill.background()
    destino = pathlib.Path(destino)
    apresentacao.save(destino)
    return destino


# -- Excel ----------------------------------------------------------------------
_NUMERO_BR = re.compile(r"^-?(\d{1,3}(\.\d{3})+|\d+)(,\d+)?$")


def _valor_celula(texto):
    if texto is None:
        return None
    texto = str(texto).strip()
    limpo = texto.replace("R$", "").replace("%", "").strip()
    if _NUMERO_BR.match(limpo):
        numero = float(limpo.replace(".", "").replace(",", "."))
        return int(numero) if numero.is_integer() and "," not in limpo \
            else numero
    return texto


def tabelas(d: Documento) -> list[tuple[int, list[list]]]:
    """(pagina, linhas) de cada tabela detectada no documento."""
    achadas = []
    for i in range(d.paginas):
        for t in d.doc[i].find_tables().tables:
            linhas = t.extract()
            if linhas and any(any(c for c in l) for l in linhas):
                achadas.append((i, linhas))
    return achadas


def para_excel(d: Documento, destino: str | os.PathLike) -> int:
    """Grava as tabelas em abas. Devolve quantas tabelas foram encontradas
    (0 = nenhuma, e o arquivo nao e' criado)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    achadas = tabelas(d)
    if not achadas:
        return 0
    livro = Workbook()
    livro.remove(livro.active)
    for n, (pagina, linhas) in enumerate(achadas, 1):
        aba = livro.create_sheet(f"Pág {pagina + 1} - T{n}"[:31])
        for linha in linhas:
            aba.append([_valor_celula(c) for c in linha])
        for celula in aba[1]:
            celula.font = Font(bold=True)
        for coluna in aba.columns:
            largura = max((len(str(c.value)) for c in coluna
                           if c.value is not None), default=8)
            aba.column_dimensions[coluna[0].column_letter].width = \
                min(60, largura + 2)
    livro.save(pathlib.Path(destino))
    return len(achadas)


# -- Markdown ---------------------------------------------------------------------
def _tamanho_corpo(doc: pymupdf.Document) -> float:
    """O tamanho de fonte com MAIS caracteres e' o do corpo do texto."""
    contagem: collections.Counter = collections.Counter()
    for p in doc:
        for b in p.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    contagem[round(s["size"])] += len(s["text"].strip())
    return float(contagem.most_common(1)[0][0]) if contagem else 11.0


_MARCADOR_LISTA = re.compile(r"^\s*([•·▪◦‣\-–*]|\d{1,3}[.)])\s+")


def _escapar(texto: str) -> str:
    return re.sub(r"([\\`*_\[\]#|])", r"\\\1", texto)


def para_markdown(d: Documento) -> str:
    doc = d.doc
    corpo = _tamanho_corpo(doc)
    partes: list[str] = []
    for i in range(doc.page_count):
        p = doc[i]
        if i:
            partes.append("\n---\n")
        areas_tabela = []
        for t in p.find_tables().tables:
            areas_tabela.append((pymupdf.Rect(t.bbox), t.extract()))
        tabelas_escritas = set()
        for b in p.get_text("dict", sort=True)["blocks"]:
            if b["type"] != 0:
                continue
            caixa = pymupdf.Rect(b["bbox"])
            dentro = [n for n, (r, _l) in enumerate(areas_tabela)
                      if r.contains(caixa) or r.intersects(caixa)
                      and (r & caixa).get_area() > caixa.get_area() * 0.6]
            if dentro:
                n = dentro[0]
                if n not in tabelas_escritas:
                    tabelas_escritas.add(n)
                    partes.append(_tabela_md(areas_tabela[n][1]))
                continue
            linhas = []
            maior = 0.0
            for l in b["lines"]:
                pedacos = []
                for s in l["spans"]:
                    t = s["text"]
                    if not t.strip():
                        pedacos.append(t)
                        continue
                    maior = max(maior, s["size"])
                    negrito = bool(s["flags"] & 16) or "bold" in s["font"].lower()
                    italico = bool(s["flags"] & 2) or "italic" in s["font"].lower()
                    conteudo = _escapar(t.strip())
                    if negrito and italico:
                        conteudo = f"***{conteudo}***"
                    elif negrito:
                        conteudo = f"**{conteudo}**"
                    elif italico:
                        conteudo = f"*{conteudo}*"
                    espaco_antes = " " if t[:1].isspace() else ""
                    espaco_depois = " " if t[-1:].isspace() else ""
                    pedacos.append(espaco_antes + conteudo + espaco_depois)
                linhas.append("".join(pedacos).strip())
            texto = " ".join(x for x in linhas if x)
            if not texto:
                continue
            razao = maior / corpo if corpo else 1
            sem_marcas = texto.replace("**", "").replace("*", "")
            if razao >= 1.6 and len(sem_marcas) < 120:
                partes.append(f"\n# {sem_marcas}\n")
            elif razao >= 1.3 and len(sem_marcas) < 140:
                partes.append(f"\n## {sem_marcas}\n")
            elif razao >= 1.12 and len(sem_marcas) < 160:
                partes.append(f"\n### {sem_marcas}\n")
            else:
                itens = [x for x in linhas if x]
                if all(_MARCADOR_LISTA.match(x.replace("\\", "")) for x in itens):
                    for x in itens:
                        bruto = x.replace("\\", "")
                        m = _MARCADOR_LISTA.match(bruto)
                        numerado = m.group(1)[0].isdigit()
                        resto = bruto[m.end():]
                        partes.append(f"{m.group(1) if numerado else '-'} {resto}")
                    partes.append("")
                else:
                    partes.append(texto + "\n")
    saida = "\n".join(partes)
    return re.sub(r"\n{3,}", "\n\n", saida).strip() + "\n"


def _tabela_md(linhas: list[list]) -> str:
    def celula(c):
        return _escapar(str(c or "").replace("\n", " ").strip())
    if not linhas:
        return ""
    largura = max(len(l) for l in linhas)
    linhas = [list(l) + [""] * (largura - len(l)) for l in linhas]
    saida = ["| " + " | ".join(celula(c) for c in linhas[0]) + " |",
             "|" + "---|" * largura]
    saida += ["| " + " | ".join(celula(c) for c in l) + " |"
              for l in linhas[1:]]
    return "\n" + "\n".join(saida) + "\n"


# -- extrair imagens --------------------------------------------------------------
def extrair_imagens(d: Documento, pasta: str | os.PathLike, *,
                    tamanho_minimo: int = 16) -> list[pathlib.Path]:
    """Salva cada imagem embutida UMA vez (a mesma logo repetida em todas as
    paginas e' um so' objeto), no formato original quando possivel. Imagens
    menores que `tamanho_minimo` px (fios, pontos) ficam de fora."""
    pasta = pathlib.Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    base = pathlib.Path(d.nome).stem
    vistas: set[int] = set()
    saidas = []
    for i in range(d.paginas):
        for n, info in enumerate(d.doc[i].get_images(full=True), 1):
            xref = info[0]
            if xref in vistas:
                continue
            vistas.add(xref)
            try:
                img = d.doc.extract_image(xref)
            except Exception:                     # noqa: BLE001
                continue
            if not img or min(img.get("width", 0),
                              img.get("height", 0)) < tamanho_minimo:
                continue
            dados, ext = img["image"], img["ext"]
            if info[1]:                     # tem mascara (transparencia)
                try:
                    pix = pymupdf.Pixmap(d.doc, xref)
                    mascara = pymupdf.Pixmap(d.doc, info[1])
                    pix = pymupdf.Pixmap(pix, mascara)
                    dados, ext = pix.tobytes("png"), "png"
                except Exception:                 # noqa: BLE001
                    pass
            if ext not in ("png", "jpeg", "jpg", "jpx", "bmp", "gif", "tiff",
                           "jp2"):
                pix = pymupdf.Pixmap(d.doc, xref)
                if pix.n - pix.alpha >= 4:          # CMYK
                    pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                dados, ext = pix.tobytes("png"), "png"
            destino = pasta / f"{base}_p{i + 1:03d}_{n}.{'jpg' if ext == 'jpeg' else ext}"
            destino.write_bytes(dados)
            saidas.append(destino)
    return saidas
