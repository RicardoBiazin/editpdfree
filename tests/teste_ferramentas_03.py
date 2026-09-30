"""Recursos da 0.3: reparar, PDF->PowerPoint/Excel/Markdown, extrair imagens,
PDF/A, marca d'agua de imagem, detectar campos e digitalizar (sem scanner)."""

from __future__ import annotations

import sys

from ajudantes import (checa, checa_igual, checa_levanta, pasta_temporaria,
                       pdf_texto, resumir, secao, texto_da_pagina)

import pymupdf

from editpdfree import (conversao_saida, digitalizar, extras, formularios,
                        pdfa, reparar)
from editpdfree.documento import Documento


def pdf_com_tabela() -> bytes:
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 60), "Relatório de vendas", fontsize=22)
    p.insert_text((72, 84), "Texto do corpo em tamanho normal.", fontsize=11)
    x0, y0, w, h = 72, 100, 120, 24
    for r in range(4):
        p.draw_line((x0, y0 + r * h), (x0 + 3 * w, y0 + r * h))
    for c in range(4):
        p.draw_line((x0 + c * w, y0), (x0 + c * w, y0 + 3 * h))
    dados = [["Produto", "Qtd", "Preço"], ["Caneta", "10", "1.234,50"],
             ["Lápis", "5", "1,20"]]
    for r, linha in enumerate(dados):
        for c, t in enumerate(linha):
            p.insert_text((x0 + c * w + 4, y0 + r * h + 16), t, fontsize=11)
    p.insert_text((72, 230), "Subtítulo da seção", fontsize=15)
    p.insert_text((72, 254), "• primeiro item", fontsize=11)
    p.insert_text((72, 270), "• segundo item", fontsize=11)
    return doc.tobytes()


def teste_reparar() -> None:
    secao("reparar")
    bom = pdf_texto([["Página um"], ["Página dois"], ["Página três"]])
    dados, rel = reparar.reparar(bom)
    checa(not rel.precisou_reparo, "arquivo bom: nao precisou de reparo")
    checa_igual(rel.paginas, 3, "3 paginas")
    # Estraga a tabela xref e o trailer
    i = bom.rfind(b"xref")
    estragado = bom[:i] + b"lixo lixo lixo\n%%EOF"
    dados, rel = reparar.reparar(estragado)
    checa(rel.precisou_reparo, "xref destruida: reparo detectado")
    with pymupdf.open("pdf", dados) as doc:
        checa_igual(doc.page_count, 3, "todas as paginas recuperadas")
        checa("Página três" in doc[2].get_text(), "conteudo preservado")
        checa(not doc.is_repaired, "o arquivo gravado abre sem reparo")
    checa("Recuperada(s) 3 de 3" in rel.resumo, "resumo legivel")
    checa_levanta(ValueError, reparar.reparar, "lixo total e' recusado",
                  b"isto nunca foi um pdf")


def teste_conversoes() -> None:
    secao("PDF para PowerPoint, Excel e Markdown")
    d = Documento(dados=pdf_com_tabela(), nome="vendas.pdf")
    with pasta_temporaria() as pasta:
        from pptx import Presentation
        saida = conversao_saida.para_powerpoint(d, pasta / "v.pptx", dpi=60)
        apres = Presentation(saida)
        checa_igual(len(apres.slides), 1, "um slide por pagina")
        textos = [s.text_frame.text for s in apres.slides[0].shapes
                  if s.has_text_frame]
        checa(any("Relatório de vendas" in t for t in textos),
              "texto editavel no slide")
        checa(any(s.shape_type == 13 for s in apres.slides[0].shapes),
              "imagem da pagina como fundo")
        checa_igual(apres.slide_width, int(595 * 12700), "slide do tamanho da pagina")

        from openpyxl import load_workbook
        n = conversao_saida.para_excel(d, pasta / "v.xlsx")
        checa_igual(n, 1, "uma tabela encontrada")
        aba = load_workbook(pasta / "v.xlsx").active
        checa_igual([c.value for c in aba[1]], ["Produto", "Qtd", "Preço"],
                    "cabecalho")
        checa_igual(aba["C2"].value, 1234.5, "1.234,50 vira numero")
        checa_igual(aba["B3"].value, 5, "inteiro vira inteiro")
        checa(aba["A1"].font.bold, "cabecalho em negrito")
        sem = Documento(dados=pdf_texto([["sem tabela"]]))
        checa_igual(conversao_saida.para_excel(sem, pasta / "x.xlsx"), 0,
                    "sem tabela: 0")
        checa(not (pasta / "x.xlsx").exists(), "e nao cria arquivo vazio")

    md = conversao_saida.para_markdown(d)
    checa("# Relatório de vendas" in md, "titulo pelo tamanho da fonte")
    checa("## Subtítulo da seção" in md or "### Subtítulo da seção" in md,
          "subtitulo vira titulo de nivel menor")
    checa("| Produto | Qtd | Preço |" in md, "tabela em Markdown")
    checa("- primeiro item" in md and "- segundo item" in md, "lista")
    checa(md.count("Caneta") == 1, "texto da tabela nao duplicado")
    dois = Documento(dados=pdf_texto([["A"], ["B"]]))
    checa("---" in conversao_saida.para_markdown(dois), "separador de pagina")


def teste_imagens() -> None:
    secao("extrair imagens")
    doc = pymupdf.open()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 40), False)
    pix.clear_with(90)
    png = pix.tobytes("png")
    for _ in range(3):                 # mesma imagem em 3 paginas
        p = doc.new_page()
        p.insert_image(pymupdf.Rect(50, 50, 200, 150), stream=png)
    p.insert_image(pymupdf.Rect(300, 50, 302, 52),
                   pixmap=pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 4, 4), False))
    d = Documento(dados=doc.tobytes(), nome="fotos.pdf")
    with pasta_temporaria() as pasta:
        saidas = conversao_saida.extrair_imagens(d, pasta)
        checa_igual(len(saidas), 1,
                    "imagem repetida sai uma vez; a minuscula fica de fora")
        img = pymupdf.Pixmap(str(saidas[0]))
        checa_igual((img.width, img.height), (60, 40),
                    "tamanho original em pixels")


def teste_pdfa() -> None:
    secao("PDF/A")
    d = Documento(dados=pdf_texto([["Contrato para arquivar"]]))
    d.doc.set_metadata({"title": "Contrato", "author": "Fulano"})
    antes = pdfa.verificar(d.doc)
    checa(any("OutputIntent" in x for x in antes), "sem OutputIntent antes")
    restantes = pdfa.para_pdfa(d)
    doc = d.doc
    checa(doc.xref_get_key(doc.pdf_catalog(), "OutputIntents")[0] == "array",
          "OutputIntent incluido")
    xmp = doc.get_xml_metadata()
    checa("<pdfaid:part>2</pdfaid:part>" in xmp and
          "<pdfaid:conformance>B</pdfaid:conformance>" in xmp, "XMP PDF/A-2b")
    checa(">Contrato<" in xmp and ">Fulano<" in xmp,
          "XMP coerente com o dicionario Info")
    checa(any("Fontes não embutidas" in x and "Helvetica" in x
              for x in restantes),
          "pendencia honesta: Helvetica nao embutida")
    checa(not any("OutputIntent" in x or "XMP" in x for x in restantes),
          "o que da' para corrigir foi corrigido")
    with pasta_temporaria() as pasta:
        d.salvar(pasta / "a.pdf")
        with pymupdf.open(pasta / "a.pdf") as c:
            checa("pdfaid:part" in c.get_xml_metadata(),
                  "XMP sobrevive ao salvar")
    # JavaScript removido
    js = pymupdf.open("pdf", pdf_texto([["x"]]))
    cat = js.pdf_catalog()
    xref = js.get_new_xref()
    js.update_object(xref, "<< /S /JavaScript /JS (app.alert(1)) >>")
    js.xref_set_key(cat, "OpenAction", f"{xref} 0 R")
    dj = Documento(dados=js.tobytes())
    checa(any("OpenAction" in x for x in pdfa.verificar(dj.doc)),
          "acao de abertura detectada")
    pdfa.para_pdfa(dj)
    checa(not any("OpenAction" in x for x in pdfa.verificar(dj.doc)),
          "acao de abertura removida")


def teste_marca_imagem() -> None:
    secao("marca d'agua de imagem")
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 50), False)
    pix.clear_with(0)
    d = Documento(dados=pdf_texto([["Texto"], ["Outra"]], rotacoes=[0, 90]))
    extras.marca_dagua_imagem(d, pix.tobytes("png"), opacidade=0.3)
    p = d.doc[0]
    info = p.get_image_info()
    checa_igual(len(info), 1, "uma imagem centralizada")
    caixa = pymupdf.Rect(info[0]["bbox"])
    centro = (caixa.tl + caixa.br) / 2
    checa(abs(centro.x - p.rect.width / 2) < 1 and
          abs(centro.y - p.rect.height / 2) < 1, "centralizada")
    render = p.get_pixmap(alpha=False)
    cor = render.pixel(int(centro.x), int(centro.y))
    checa(150 < cor[0] < 220, f"translucida (pixel {cor}, preto a 30%)")
    p1 = d.doc[1]
    caixa1 = pymupdf.Rect(p1.get_image_info()[0]["bbox"]) * p1.rotation_matrix
    caixa1.normalize()
    checa(caixa1.width > caixa1.height,
          "pagina girada: a imagem continua deitada para o leitor")
    d2 = Documento(dados=pdf_texto([["x"]]))
    extras.marca_dagua_imagem(d2, pix.tobytes("png"), escala=0.2,
                              lado_a_lado=True)
    p2 = d2.doc[0]
    checa(len(p2.get_image_info()) > 4, "lado a lado repete a imagem")
    checa_levanta(ValueError, extras.marca_dagua_imagem, "opacidade invalida",
                  d2, pix.tobytes("png"), opacidade=0)


def teste_detectar_campos() -> None:
    secao("detectar campos")
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 100), "Nome: ______________________", fontsize=12)
    p.insert_text((72, 140), "CPF:", fontsize=12)
    p.draw_line((110, 142), (300, 142))
    p.draw_rect(pymupdf.Rect(72, 170, 84, 182))
    p.insert_text((90, 180), "Aceito os termos", fontsize=12)
    p.draw_rect(pymupdf.Rect(72, 300, 400, 500))            # moldura grande
    d = Documento(dados=doc.tobytes())
    sug = formularios.detectar_campos(d)
    tipos = sorted((s.tipo, s.nome) for s in sug)
    checa_igual(tipos, [("caixa", "aceito_os_termos"), ("texto", "cpf"),
                        ("texto", "nome")],
                "sublinhado, linha e quadrado; a moldura grande nao conta")
    n = formularios.criar_sugeridos(d, sug)
    checa_igual(n, 3, "tres campos criados")
    checa_igual(len(formularios.campos(d)), 3, "campos existem no PDF")
    checa_igual(formularios.detectar_campos(d), [],
                "depois de criados nao sao sugeridos de novo")


def teste_digitalizar() -> None:
    secao("digitalizar (sem scanner)")
    imagens = []
    for cinza in (200, 120):
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 2480, 3508), False)
        pix.clear_with(cinza)
        pix.set_dpi(300, 300)
        imagens.append(pix.tobytes("jpg"))
    dados = digitalizar.montar_pdf(imagens)
    with pymupdf.open("pdf", dados) as doc:
        checa_igual(doc.page_count, 2, "uma pagina por imagem")
        r = doc[0].rect
        checa(abs(r.width - 595.2) < 1 and abs(r.height - 841.9) < 1,
              f"tamanho fisico pelo DPI (A4 a 300 dpi: {r.width:.1f})")
    checa_levanta(ValueError, digitalizar.montar_pdf, "sem imagens", [])
    checa(isinstance(digitalizar.scanners(), list),
          "listar scanners nao abre dialogo nem falha")


def main() -> int:
    teste_reparar()
    teste_conversoes()
    teste_imagens()
    teste_pdfa()
    teste_marca_imagem()
    teste_detectar_campos()
    teste_digitalizar()
    return resumir()


if __name__ == "__main__":
    sys.exit(main())
