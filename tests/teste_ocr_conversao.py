"""OCR (Tesseract embutido no MuPDF) e conversoes de/para PDF."""

from __future__ import annotations

import sys

from ajudantes import (checa, checa_igual, checa_levanta, pasta_temporaria,
                       pdf_texto, resumir, secao)

import pymupdf

from editpdfree import conversao, ocr, seguranca
from editpdfree.documento import Documento


def escaneado(linhas: list[str], rotacao: int = 0) -> bytes:
    """PDF so' com IMAGEM do texto -- como um escaneado."""
    src = pymupdf.open()
    p = src.new_page()
    for i, linha in enumerate(linhas):
        p.insert_text((72, 120 + 30 * i), linha, fontsize=16)
    pix = p.get_pixmap(dpi=200)
    scan = pymupdf.open()
    sp = scan.new_page(width=p.rect.width, height=p.rect.height)
    sp.insert_image(sp.rect, pixmap=pix)
    if rotacao:
        sp.set_rotation(rotacao)
    return scan.tobytes()


def teste_ocr() -> None:
    secao("OCR")
    d = Documento(dados=escaneado(["Contrato de locação residencial",
                                   "Florianópolis, Santa Catarina"]))
    checa_igual(d.doc[0].get_text().strip(), "", "escaneado nao tem texto")
    tamanho_antes = len(d.para_bytes())
    feitas = ocr.reconhecer(d, idioma="por")
    checa_igual(feitas, 1, "uma pagina reconhecida")
    texto = d.doc[0].get_text()
    checa("locação" in texto, "reconhece portugues com acento")
    checa("Florianópolis" in texto, "segunda linha")
    caixa = d.doc[0].search_for("Florianópolis")
    checa(caixa and abs(caixa[0].x0 - 72) < 6 and abs(caixa[0].y1 - 150) < 8,
          f"texto reconhecido na posicao certa ({caixa[:1]})")
    checa(len(d.para_bytes()) < tamanho_antes * 1.3,
          "a camada de OCR nao duplica a imagem no arquivo")
    n = len(d._desfazer)
    checa_igual(ocr.reconhecer(d), 0, "pagina com texto e' pulada")
    checa_igual(len(d._desfazer), n, "nada a fazer nao cria desfazer")
    # Tarjar funciona sobre o texto reconhecido
    seguranca.tarjar_texto(d, "Florianópolis")
    checa("Florianópolis" not in d.doc[0].get_text(),
          "tarjar o texto reconhecido remove a camada tambem")
    d.desfazer()
    d.desfazer()
    checa_igual(d.doc[0].get_text().strip(), "", "desfazer tira a camada")

    secao("OCR em pagina girada")
    g = Documento(dados=escaneado(["Documento girado noventa"], rotacao=90))
    ocr.reconhecer(g, idioma="por")
    checa("girado" in g.doc[0].get_text(), "reconhece pagina girada")
    checa_igual(g.doc[0].rotation, 90, "a rotacao da pagina e' preservada")
    caixa = g.doc[0].search_for("girado")
    checa(caixa and caixa[0].y1 < 140 and caixa[0].x0 > 72,
          "posicao no espaco sem rotacao")

    secao("cancelar")
    c = Documento(dados=escaneado(["A"]))
    checa_igual(ocr.reconhecer(c, progresso=lambda i, t: False), 0,
                "cancelar antes da primeira pagina")


def teste_conversao() -> None:
    secao("para PDF")
    with pasta_temporaria() as pasta:
        import docx
        w = docx.Document()
        w.add_heading("Relatório anual", 1)
        w.add_paragraph("Texto com acentuação e ç.")
        arquivo = pasta / "rel.docx"
        w.save(arquivo)
        dados = conversao.para_pdf(arquivo, usar_libreoffice=False)
        with pymupdf.open("pdf", dados) as doc:
            texto = doc[0].get_text()
            checa("Relatório anual" in texto, "DOCX -> PDF (MuPDF)")
            checa("acentuação" in texto, "acentos preservados")
        txt = pasta / "nota.txt"
        txt.write_text("Linha um\nLinha dois ação", "utf-8")
        with pymupdf.open("pdf", conversao.para_pdf(txt)) as doc:
            checa("Linha dois ação" in doc[0].get_text(), "TXT -> PDF")
        img = pasta / "i.png"
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 80), False)
        pix.clear_with(128)
        pix.save(img)
        with pymupdf.open("pdf", conversao.para_pdf(img)) as doc:
            checa_igual(doc.page_count, 1, "imagem -> PDF")
        checa_levanta(ValueError, conversao.para_pdf, "formato desconhecido",
                      pasta / "x.xyz")
        junto = conversao.varios_para_pdf([arquivo, txt, img],
                                          pasta / "junto.pdf")
        with pymupdf.open(junto) as doc:
            checa(doc.page_count >= 3, "varios formatos num PDF so'")

        secao("PDF para Word")
        d = Documento(dados=pdf_texto([["Título do contrato",
                                        "Cláusula primeira: objeto"],
                                       ["Segunda página"]]))
        saida = conversao.para_word(d, pasta / "c.docx")
        lido = docx.Document(saida)
        texto = "\n".join(p.text for p in lido.paragraphs)
        checa("Título do contrato" in texto, "texto no Word")
        checa("Cláusula primeira: objeto" in texto, "acentos no Word")
        checa("Segunda página" in texto, "segunda pagina")
        quebras = sum(1 for p in lido.paragraphs for r in p.runs
                      if 'w:br' in r._r.xml and 'type="page"' in r._r.xml)
        checa_igual(quebras, 1, "quebra de pagina entre as paginas")
        run = next(r for p in lido.paragraphs for r in p.runs
                   if "Título" in r.text)
        checa(run.font.size and abs(run.font.size.pt - 12) < 0.6,
              "tamanho da fonte preservado")


def main() -> int:
    teste_ocr()
    teste_conversao()
    return resumir()


if __name__ == "__main__":
    sys.exit(main())
